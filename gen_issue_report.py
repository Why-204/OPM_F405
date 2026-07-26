# -*- coding: utf-8 -*-
"""生成《OPM_F405 光功率计协议 已知问题清单》Word 文档。"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import datetime

OUT = r"D:\hardware\lightpowerconsume\OPM_F405_已知问题清单.docx"

doc = Document()

# 中文字体
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_cn(run, name="Microsoft YaHei"):
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_cn(run)
    return p


def para(text, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    set_cn(run)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        r0.bold = True
        set_cn(r0)
    r = p.add_run(text)
    set_cn(r)
    return p


def code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    shd = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear")
    e.set(qn("w:fill"), "F2F2F2")
    shd.append(e)
    return p


RED = RGBColor(0xC0, 0x00, 0x00)
ORANGE = RGBColor(0xC0, 0x60, 0x00)
YELLOW = RGBColor(0xB0, 0x90, 0x00)

# ---------------- 封面标题 ----------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("OPM_F405 光功率计协议\n已知问题清单")
r.bold = True
r.font.size = Pt(20)
set_cn(r)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("STM32F405 / FreeRTOS / UART5(RS232) + UART3(CH9121 以太网)")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
set_cn(r)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("生成日期：%s" % datetime.date.today().isoformat())
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
set_cn(r)

doc.add_paragraph()

# ---------------- 一、概览表 ----------------
h("一、问题概览", 1)
para("下表按严重度列出已核对手册后发现、会导致功能不正确或无法实现的问题。严重度：🔴高 / 🟠中 / 🟡低。", size=10)

table = doc.add_table(rows=1, cols=5)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
heads = ["编号", "严重度", "问题", "手册依据", "现状"]
widths = [0.6, 0.8, 2.4, 1.8, 1.8]
for c, t, w in zip(hdr, heads, widths):
    c.width = Inches(w)
    p = c.paragraphs[0]
    run = p.add_run(t)
    run.bold = True
    run.font.size = Pt(9.5)
    set_cn(run)

rows = [
    ("A", "🔴 高", "RDPR 返回原始 dBmA，未减标定偏移量",
     "手册 17)「光功率=测量值−偏移量」；16) RDPR 应返回光功率 dBm",
     "get_power_dbm 直接返回 dBmA"),
    ("B", "🔴 高", "整个标定功能是死的：offset 写入/回读却从不参与计算",
     "同上", "WRPO/RDPO 存取 offset_db，RDPR 不用"),
    ("C", "🟠 中", "STWW 设的工作波长不影响任何测量，也未用于选偏移量",
     "手册 3.1 不同波长响应度不同", "work_wl 只存不用"),
    ("D", "🟠 中", "STTM 设的采样平均时间不生效",
     "手册 11) 采样时间影响精度/更新率", "sample_us 只存不用，ADC 用固定平均"),
    ("E", "🔴 高", "RDMR/连续测量为占位桩，且响应可达 ~65KB 超出帧上限",
     "手册 20~25)", "无高速采集缓冲；OPM_MAX_FRAME_LEN=300"),
    ("F", "🟡 低", "多 UART 同时重配 UART3 存在窄竞争窗口",
     "—", "已标记，未处理"),
]
for num, sev, prob, basis, cur in rows:
    cells = table.add_row().cells
    vals = [num, sev, prob, basis, cur]
    for i, (c, t) in enumerate(zip(cells, vals)):
        c.width = Inches(widths[i])
        p = c.paragraphs[0]
        run = p.add_run(t)
        run.font.size = Pt(9)
        if i == 1:
            if "高" in t:
                run.font.color.rgb = RED
            elif "中" in t:
                run.font.color.rgb = ORANGE
            else:
                run.font.color.rgb = YELLOW
            run.bold = True
        set_cn(run)

doc.add_paragraph()

# ---------------- 二、核心问题链 ----------------
h("二、核心问题：A + B + C 是一根断裂的链条", 1)
para("手册的测量模型是：", size=10.5)
code("光功率(dBm) = 测量值(dBmA) − offset[通道][当前工作波长]")
para("目前这条链的三个环节全部断开：", size=10.5)
bullet("RDPR 不减 offset。", "1) ")
bullet("offset 表 WRPO 写进去、RDPO 读得出，但从没有人在计算里用它。", "2) ")
bullet("RDPR 不知道该用哪个波长的 offset——work_wl(nm) 没有映射到标定波长下标。", "3) ")
para("后果：", bold=True)
para("上位机完成波长标定（WRPO 写偏移量）、切换工作波长（STWW）后再读 RDPR，"
     "读到的值完全不受这些设置影响。等于标定与工作波长功能对外不可用。", size=10.5)

# ---------------- 三、逐条详情 ----------------
h("三、逐条详情", 1)

h("A / B　RDPR 未应用标定偏移量（🔴 高）", 2)
para("现状：h_rdpr 调用 get_power_dbm，后者直接返回 adc_ch[i].dBmA，未减 offset。", size=10)
para("手册 16) RDPR 返回的应是光功率 dBm；手册 17) 明确定义"
     "「光功率 = 光功率测量值 − 偏移量」。", size=10)
para("建议修复（集中在 get_power_dbm）：", bold=True, size=10)
code(
    "static float get_power_dbm(uint8_t ch_index) {\n"
    "    float raw = adc_ch[ch_index].dBmA;\n"
    "    uint16_t wl = g_opm_dev.work_wl[ch_index];   // 当前工作波长(nm)\n"
    "    for (uint8_t i = 0; i < g_opm_dev.cal_wl_count; i++)\n"
    "        if (g_opm_dev.cal_wl[i] == wl)\n"
    "            return raw - g_opm_dev.offset_db[ch_index][i];\n"
    "    return raw;   // 波长不在标定表 → 不减(或按需报错)\n"
    "}"
)

h("C　工作波长不参与测量（🟠 中）", 2)
para("work_wl 目前只被 STWW 写、RDWW 读，没有任何测量逻辑消费它。"
     "修复 A/B 后，它将用于选择对应波长的偏移量；此项随 A/B 一并解决。", size=10)

h("D　采样平均时间不生效（🟠 中）", 2)
para("sample_us 只被 STTM 写、RDTM 读，ADC 任务使用固定平均点数，"
     "通过协议修改采样时间对实际测量无效。", size=10)
para("修复需改动 ADC 任务：读取 g_opm_dev.sample_us[ch] 动态调整平均点数，"
     "改动较大，需先确认 ADC_Driver 的平均机制。", size=10)

h("E　连续测量 RDMR 未实现（🔴 高，即原 #3）", 2)
para("RDMR(25) 为占位桩，返回假数据；STMP/STMT/STST 未做真正的高速定时采集与环形缓冲。", size=10)
para("另有帧结构限制：RDMR 响应最多可达 16380 样本 × 4B ≈ 65KB，"
     "远超 OPM_MAX_FRAME_LEN=300，当前缓冲无法容纳，需重新设计分帧/传输方案。", size=10)

h("F　多 UART 并发重配 UART3（🟡 低）", 2)
para("两路 UART 若几乎同时到达命令，一个任务在重配 UART3（写 CH9121）时另一任务正在使用它，"
     "存在极窄竞争窗口。已知并已标记，暂未处理。", size=10)

# ---------------- 四、已完成/正常项 ----------------
h("四、本轮已修复 / 已确认正常的项", 1)
bullet("dBmA 计算已从 HMI 任务移到 ADC 任务数据源头（原问题 #2）。", "✔ ")
bullet("CH9121 IP/端口配置：WRIP/WRPT → netcfg_pending 标志 → 驱动重配，响应先于复位发出。", "✔ ")
bullet("开机自愈波特率：ch9121_ensure_baudrate 读-比对-写，保护 EEPROM。", "✔ ")
bullet("帧格式、RDPN/RDSN/RDVR/RDMC/RDIP/RDPT/RDCC/RDWC/RDWL/RDWW/STWW 等请求响应字节布局与手册一致。", "✔ ")
bullet("RDPR/RDPO/WRPO 请求中通道号后固定 0x01 标志字节已正确处理与回显。", "✔ ")
bullet("中断优先级 USART3/UART5/UART4/DMA1_Stream0/1 均为 5，ISR 内 osThreadFlagsSet 安全。", "✔ ")

# ---------------- 五、待你确认 ----------------
h("五、需你确认的语义点", 1)
para("修复 A 的前提：adc_ch[i].dBmA 是否就是手册所说的『测量值』，即"
     "「光功率 = dBmA − offset」这个减法是否成立？", size=10.5)
bullet("若是（设备内部即 dBmA，靠 offset 转真实 dBm）→ 上面的修复直接可用。", "情况一：")
bullet("若 dBmA 与 dBm 间还有响应度(A/W)等换算 → 需先厘清换算关系再改。", "情况二：")

doc.add_paragraph()
para("确认后可一次性修复 A+B+C；D 是否一并做取决于是否需要采样时间生效（需先看 ADC_Driver）。",
     bold=True, size=10)

doc.save(OUT)
print("SAVED", OUT)
