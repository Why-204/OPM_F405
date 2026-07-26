from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path(r"D:\quantization\量化交易与量化投资学习手册.docx")
BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GOLD = "7A5A00"
RED = "9B1C1C"
GRAY = "666666"
INK = "1F2937"
FONT_CN = "Microsoft YaHei"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

def font(run, size=11, bold=None, color=INK, italic=None):
    run.font.name = FONT_CN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    return run

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT_CN
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [
    ("Title", 30, DARK, 0, 10), ("Subtitle", 14, GRAY, 0, 18),
    ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK, 10, 5)]:
    s = styles[name]
    s.font.name = FONT_CN
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = name != "Subtitle"
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
for list_name in ["List Bullet", "List Number"]:
    s = styles[list_name]
    s.font.name = FONT_CN
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    s.font.size = Pt(11)
    s.paragraph_format.left_indent = Inches(.375)
    s.paragraph_format.first_line_indent = Inches(-.188)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.25

if "Callout" not in styles:
    callout_style = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
else:
    callout_style = styles["Callout"]
callout_style.font.name = FONT_CN
callout_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
callout_style.font.size = Pt(10.5)
callout_style.paragraph_format.space_before = Pt(6)
callout_style.paragraph_format.space_after = Pt(8)
callout_style.paragraph_format.left_indent = Inches(.16)
callout_style.paragraph_format.right_indent = Inches(.16)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc.get_or_add_tcPr()
    tcMar = tc.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tc.append(tcMar)
    for tag, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn("w:"+tag))
        if node is None: node = OxmlElement("w:"+tag); tcMar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None: tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total)); tblW.set(qn("w:type"), "dxa")
    ind = tblPr.find(qn("w:tblInd"))
    if ind is None: ind = OxmlElement("w:tblInd"); tblPr.append(ind)
    ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for ch in list(grid): grid.remove(ch)
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(w)); grid.append(gc)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None: tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[idx])); tcW.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx]/1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    if widths is None: widths = [9360//len(headers)]*len(headers)
    set_table_geometry(t, widths)
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], LIGHT)
        p = t.rows[0].cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(str(h)), 10, True, DARK)
    set_repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
            if i == 0 and len(headers) <= 3: font(p.add_run(str(val)), 9.5, True, DARK)
            else: font(p.add_run(str(val)), 9.5)
    set_table_geometry(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def h1(x): doc.add_heading(x, level=1)
def h2(x): doc.add_heading(x, level=2)
def h3(x): doc.add_heading(x, level=3)
def para(x, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and x.startswith(bold_prefix):
        font(p.add_run(bold_prefix), 11, True, DARK)
        font(p.add_run(x[len(bold_prefix):]), 11)
    else: font(p.add_run(x), 11)
    return p
def bullet(x, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    font(p.add_run(x), 11); return p
def num(x):
    p = doc.add_paragraph(style="List Number"); font(p.add_run(x), 11); return p
def callout(label, text, color=DARK):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; set_table_geometry(t, [9360])
    shade(t.cell(0,0), PALE); p=t.cell(0,0).paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    font(p.add_run(label+"："), 10.5, True, color); font(p.add_run(text), 10.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def page_break(): doc.add_page_break()

# Header / footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("量化学习路线与成果验收手册"), 9, False, GRAY)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("仅用于学习与研究，不构成投资建议"), 9, False, GRAY)

# Cover
for _ in range(5): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("量化交易与量化投资"), 30, True, DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("系统学习手册"), 24, True, BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("24周路线 · 学习内容 · 项目实践 · 成果验收"), 14, False, GRAY)
doc.add_paragraph()
callout("核心目标", "从零开始建立一套可重复、可审计的量化研究流程：提出假设、获取数据、实现策略、可信回测、风险评估、样本外验证和模拟运行。")
for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("建议投入：每周 8～12 小时｜适用方向：股票、ETF、中低频策略"), 10.5, False, GRAY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("版本：2026年7月"), 10, False, GRAY)
page_break()

h1("使用说明")
para("这不是一份“看完就会”的课程目录，而是一份以交付物为中心的学习计划。每个阶段都要求你留下代码、研究记录、图表或报告，并用明确标准判断自己是否真正掌握。")
callout("学习原则", "先保证回测正确，再追求收益；先用简单模型建立基线，再增加复杂度；先模拟运行，再考虑真实资金。")
h2("推荐执行方式")
for x in ["固定每周至少4次学习时段，每次60～120分钟。","每学一个概念，必须完成一个最小代码例子。","每个策略必须保留研究假设、数据版本、参数、成本假设和结果。","每周复盘一次，每四周做一次阶段验收。","验收未通过时，不进入下一阶段，先完成补救任务。"]: bullet(x)
h2("全程最终成果")
for x in ["一个结构清晰的量化研究代码仓库。","一套行情数据清洗与质量检查程序。","一个无明显未来函数、包含费用与滑点的回测器。","至少三种策略的研究报告，其中一项完成样本外和模拟交易验证。","一份个人量化研究规范与风险纪律。"]: bullet(x)
h2("路线总览")
table(["阶段","周数","核心能力","必须交付"],[
("0. 定位与环境","第1周","理解量化工作流，建立环境","环境截图、学习日志模板"),
("1. Python数据基础","第2～5周","处理表格与时间序列","数据分析Notebook"),
("2. 金融与统计基础","第6～8周","正确计算收益和风险","绩效分析器"),
("3. 手写回测","第9～12周","理解信号、持仓、成交","双均线完整回测"),
("4. 可信回测","第13～16周","识别偏差和过拟合","回测审计报告"),
("5. 策略拓展","第17～20周","研究趋势、反转、配置","三策略对比报告"),
("6. 模拟运行","第21～24周","研究工程化与纪律","模拟交易月报"),
], [1300,1500,2500,4060])

h1("第0阶段：明确方向与搭建环境（第1周）")
h2("学习目标")
for x in ["区分量化研究、回测、模拟交易和实盘交易。","了解中低频量化、统计套利、高频交易和资产配置的差异。","选择一个起步市场和一个可控的研究问题。","搭建可重复运行的Python研究环境。"]: bullet(x)
h2("学习内容")
h3("0.1 认识量化研究闭环")
para("标准闭环是：提出可证伪假设 → 定义数据与交易规则 → 编写代码 → 回测 → 检查偏差 → 样本外验证 → 模拟运行 → 复盘。任何无法写成精确规则的“感觉”都还不是可执行策略。")
h3("0.2 工具环境")
for x in ["Python 3、VS Code或PyCharm、Jupyter Notebook。","NumPy：数组和向量化计算。","pandas：表格、时间索引、滚动窗口和数据清洗。","Matplotlib/Seaborn：价格、净值、回撤和分布图。","Git：记录策略与数据处理代码的修改历史。"]: bullet(x)
h3("0.3 建议目录")
table(["目录","用途"],[
("data/raw","原始数据，只读保存，不直接修改"),
("data/processed","清洗、复权、对齐后的数据"),
("notebooks","探索分析和教学练习"),
("src","可复用的数据、指标、回测、绩效代码"),
("strategies","策略规则和参数"),
("reports","图表、回测报告、复盘"),
("tests","单元测试与防未来函数测试"),
],[1900,7460])
h2("第1周成果验收")
table(["检查项","通过标准","补救动作"],[
("环境","能从头创建环境并运行Notebook","重新安装并写一页安装记录"),
("数据","能读入CSV并显示前5行、列类型、日期范围","练习read_csv、info、describe"),
("认知","能用自己的话说明研究、回测、模拟、实盘的区别","录制3分钟口述或写300字"),
("管理","建立Git仓库并完成至少3次有意义提交","学习add/commit/log并重新操作"),
],[1800,4200,3360])

h1("第1阶段：Python与时间序列数据（第2～5周）")
h2("第2周：Python核心语法")
for x in ["变量、数字、字符串、布尔值和None。","列表、元组、字典、集合。","if/elif/else、for、while。","函数参数、返回值、作用域与异常处理。","模块导入、虚拟环境和包。"]: bullet(x)
h3("练习")
for x in ["编写函数计算单笔交易收益率。","输入一组日收益率，计算累计净值。","读取交易记录列表，统计盈利和亏损笔数。","遇到除零、空数据和非法价格时给出明确错误。"]: num(x)
h3("验收")
callout("通过标准", "不查答案，45分钟内编写一个函数：输入价格列表，返回日收益、累计收益、最大单日涨跌，并正确处理缺失值。至少写5个测试案例。")

h2("第3周：NumPy")
for x in ["ndarray、shape、dtype、索引、切片。","布尔筛选、广播、聚合、向量化。","nan处理、随机数和可复现实验。","理解视图与复制，避免修改原数据。"]: bullet(x)
h3("验收任务")
para("生成1000个模拟日收益率，用向量化方法计算净值、年化收益、波动率和最大回撤；不得使用逐行for循环完成核心计算。")
table(["评分项","分值","达标描述"],[
("结果正确",40,"与手工小样本结果一致"),
("向量化",25,"核心计算无逐行循环"),
("边界处理",20,"空数组、NaN、全负收益可处理"),
("可复现",15,"固定随机种子，重复运行一致"),
],[1800,900,6660])

h2("第4周：pandas")
for x in ["Series、DataFrame、索引与列选择。","CSV/Excel读写、类型转换、缺失值和重复值。","groupby、merge、concat、pivot。","DatetimeIndex、resample、shift、rolling、expanding。"]: bullet(x)
h3("关键练习")
for x in ["将日期字符串转成DatetimeIndex并排序去重。","计算1日、5日、20日收益率。","计算20日均线和20日波动率。","合并价格表与基准指数表。","找出缺失交易日、重复记录和异常价格。"]: num(x)
callout("高频错误", "把日期当普通字符串；数据未排序就rolling；不同证券交易日未对齐；用0随意填补缺失价格；忽略复权方式。", RED)

h2("第5周：可视化与探索分析")
for x in ["价格曲线、对数价格、成交量。","收益率直方图、箱线图、QQ图。","滚动波动率、滚动相关性。","净值曲线与水下回撤图。","图表标题、坐标、单位、图例和数据区间。"]: bullet(x)
h3("阶段项目A：行情体检报告")
para("选择一只宽基ETF，制作一份不少于6张图的行情体检报告。报告应包括数据来源和范围、缺失与异常检查、收益分布、年度收益、波动率、最大回撤和结论。")
table(["等级","标准"],[
("优秀","所有计算有函数封装和测试；图表标注完整；能解释异常及处理依据"),
("通过","数据正确、图表完整、结论与图表一致"),
("未通过","日期/收益计算错误，或无法说明缺失值和复权处理"),
],[1500,7860])

h1("第2阶段：金融、统计与绩效指标（第6～8周）")
h2("第6周：市场与交易机制")
for x in ["股票、ETF、指数、基金、期货的基本差异。","OHLCV含义、成交量、流动性和价差。","市价单、限价单、停牌、涨跌停。","佣金、税费、滑点、冲击成本。","前复权、后复权、不复权的用途。","A股T+1、最小交易单位和不能成交情形。"]: bullet(x)
h3("验收题")
for x in ["为什么依赖当天收盘价生成的信号通常不能假设仍以当天收盘价成交？","涨停时出现买入信号，回测器应该如何处理？","为什么回测收益要同时报告基准？","前复权数据适合研究收益，但为什么下单价格仍要谨慎处理？"]: num(x)

h2("第7周：收益与风险")
table(["指标","必须掌握","自测要求"],[
("简单/对数收益","公式、适用场景和聚合差异","手算3日样本并用代码验证"),
("年化收益","几何年化，不把平均日收益直接乘252","不同区间可正确计算"),
("年化波动率","日波动率乘√252的前提","能解释非稳定波动的限制"),
("最大回撤","峰值到后续谷值的最大跌幅","输出幅度、峰值日、谷值日"),
("夏普比率","超额收益与波动的关系","说明无风险利率和频率"),
("胜率/盈亏比","不能单独代表策略质量","用两组策略说明反例"),
],[1700,4300,3360])
h3("阶段成果：绩效分析器")
for x in ["输入日收益序列，输出累计收益、年化收益、波动率、夏普、最大回撤。","输出年度收益表、月度收益表和回撤图。","支持基准对比。","对空数据、常数收益、NaN和极端值有合理处理。"]: bullet(x)

h2("第8周：统计基础")
for x in ["均值、中位数、分位数、方差、标准差。","协方差、相关系数及相关不等于因果。","抽样误差、置信区间、假设检验和p值。","多重检验：尝试越多，偶然“有效”的策略越多。","线性回归、残差、R²和系数稳定性。","时间序列自相关、平稳性的基本概念。"]: bullet(x)
callout("验收标准", "能说明“回测夏普为2”为什么不足以证明策略有效，并至少列出数据质量、交易成本、样本长度、过拟合、未来泄漏和市场状态六类审查问题。")

h1("第3阶段：从零手写回测（第9～12周）")
h2("第9周：信号、仓位和收益")
para("必须严格区分三个时间点：何时观察数据、何时产生信号、何时能够成交。")
table(["对象","含义","典型错误"],[
("signal","模型希望持有的方向","使用未来价格生成"),
("position","实际持仓","信号生成当天立即享受全天收益"),
("trade","仓位变化","忽略换手、费用和成交约束"),
("return","持仓期间获得的收益","价格复权与现金流重复计算"),
],[1500,3500,4360])
h3("必做实验")
para("实现双均线策略，并分别比较“不shift”“shift 1期”“下一日开盘成交”的结果，解释差异来自哪里。")

h2("第10周：交易成本与成交模拟")
for x in ["双边佣金、卖出税费、最小佣金。","固定滑点、百分比滑点和买卖价差。","停牌、涨跌停、成交量不足。","仓位取整、现金不足和部分成交。"]: bullet(x)
h3("验收")
callout("通过标准", "手工设计10个交易日的小样本，逐日核对现金、持仓、成交、费用和总资产。程序结果必须与手算一致，误差只允许来自明确的舍入规则。")

h2("第11周：回测器结构")
table(["模块","职责","最低接口"],[
("Data","加载、清洗、对齐和提供历史数据","load/validate/get_bar"),
("Strategy","根据当时可见信息生成目标仓位","generate_signal"),
("Broker","检查约束、撮合、费用、现金与持仓","submit/fill"),
("Portfolio","记录资产、收益、暴露和回撤","mark_to_market"),
("Report","计算指标、图表和交易明细","evaluate/export"),
],[1500,4300,3560])
para("你的嵌入式经验可以直接迁移：行情事件类似硬件事件，策略类似业务任务，Broker类似驱动层，订单状态类似状态机。模块之间应通过明确接口通信，而不是到处修改全局变量。")

h2("第12周：阶段项目B——双均线完整回测")
for x in ["明确研究假设、资产、区间、参数和基准。","信号只使用当时已知数据。","包含费用、滑点和成交时间假设。","输出交易明细、净值、回撤和年度结果。","至少测试三个不重叠区间。","写出策略失效情形，不只写优点。"]: bullet(x)
table(["评分维度","权重","及格线"],[
("正确性","35%","无明显未来泄漏，小样本逐日核对通过"),
("真实性","20%","成本、滑点和成交时点明确"),
("代码质量","15%","模块化、无重复、关键函数有测试"),
("研究质量","20%","有基准、分期、敏感性和失败分析"),
("可复现性","10%","固定环境、参数、数据版本，可一键运行"),
],[1900,1000,6460])

h1("第4阶段：可信回测与反过拟合（第13～16周）")
h2("第13周：常见回测偏差")
table(["偏差","表现","检测方法"],[
("未来函数","用未来价格、未来成分或当日未公布数据","逐列写明信息可得时间；单步回放"),
("幸存者偏差","只使用今天仍存在的证券","使用历史成分和退市样本"),
("选择偏差","看到结果后更换起点、资产或规则","预先登记研究假设"),
("过拟合","参数微调后历史极好、换区间失效","样本外、参数热图、简化规则"),
("数据窥探","反复查看测试集并继续调参","锁定测试集，仅最终评估"),
("成本低估","高换手策略收益虚高","费用和滑点压力测试"),
],[1500,4000,3860])
h2("第14周：样本内与样本外")
for x in ["按时间顺序划分训练、验证、测试集。","测试集不参与参数选择。","使用滚动/扩展窗口进行Walk-forward验证。","报告不同牛熊、震荡和高低波动阶段。"]: bullet(x)
h3("验收")
para("选定策略后冻结代码和参数，对此前未查看的测试区间运行一次。无论结果好坏都保留报告，不允许看到结果后重新定义测试区间。")

h2("第15周：参数稳定性与压力测试")
for x in ["绘制参数网格热图，不只报告最优点。","费用提高1倍、2倍后的结果。","信号延迟1～3天后的结果。","随机删除少量交易或扰动成交价。","更换相似资产和相邻时间区间。","检查收益是否依赖极少数交易。"]: bullet(x)
callout("健康特征", "有效策略更可能在一片参数区域内表现尚可，而不是只有一个尖锐参数点异常优秀；收益来源应能被经济逻辑解释。")

h2("第16周：回测审计")
table(["审计问题","通过条件"],[
("信息时间","每个字段都注明何时可得"),
("交易时间","信号、委托、成交顺序无矛盾"),
("数据版本","来源、下载时间、复权、清洗规则可追溯"),
("交易约束","费用、滑点、停牌、涨跌停、整手有处理"),
("样本外","至少一个真正未参与调参的时间区间"),
("稳定性","参数、成本、延迟和市场分段测试完成"),
("可复现","新环境按README可以复现主要结果"),
],[2100,7260])
h3("阶段项目C")
para("对阶段项目B进行独立审计。先写审计清单，再运行测试。发现问题时保留修复前后结果，并解释收益变化。能够主动发现自己的回测错误，比得到漂亮收益更重要。")

h1("第5阶段：策略类型与组合（第17～20周）")
h2("第17周：趋势与动量")
for x in ["均线、通道突破、时间序列动量。","趋势策略在震荡市场的反复止损。","波动率调整仓位和移动止损。","不同持有周期与换手成本。"]: bullet(x)
h2("第18周：均值回归")
for x in ["价格偏离均线、布林带、短期反转。","均值是否稳定，结构变化会导致什么。","尾部风险：小赚多次、偶尔大亏。","流动性、拥挤和不能成交风险。"]: bullet(x)
h2("第19周：截面选股与因子")
for x in ["定义股票池，避免使用未来成分。","价值、动量、质量、低波动等因子的直观含义。","排序、分组、分层收益和多空组合。","市值、行业和风格暴露。","调仓频率、换手和因子衰减。"]: bullet(x)
h2("第20周：资产配置与组合")
for x in ["相关性和分散化。","等权、风险预算、波动率目标。","再平衡频率和交易成本。","组合层最大仓位、最大回撤和风险暴露限制。"]: bullet(x)
h3("阶段项目D：三策略比较")
table(["必须比较","要求"],[
("策略","趋势、反转、资产配置各一个"),
("统一条件","相同数据截止日、成本框架和报告口径"),
("市场阶段","至少按年份和高/低波动区间拆分"),
("风险","回撤、尾部损失、换手、集中度"),
("结论","说明每种策略适用与失效环境"),
],[1800,7560])

h1("第6阶段：模拟交易与研究工程化（第21～24周）")
h2("第21周：从研究代码到每日流程")
for x in ["数据更新、校验、信号生成、订单建议、日志和报告分层。","所有运行记录时间、数据版本、代码版本和参数。","发生数据缺失或异常时默认停止，而不是猜测。","保持研究环境与模拟运行环境的一致性。"]: bullet(x)
h2("第22周：模拟交易")
for x in ["每天固定时间运行。","记录理论信号、可成交价格、实际模拟成交和偏差。","不因为连续亏损临时改规则。","规则变更必须创建新版本，从新验证周期开始。"]: bullet(x)
h2("第23周：风险与故障预案")
table(["风险","预先规则示例"],[
("单资产集中","单资产目标权重不超过组合的20%"),
("组合回撤","达到预设阈值后停止新增仓位并复核"),
("数据异常","日期重复、价格跳变或数据断层时停止运行"),
("程序异常","订单数量、现金、持仓不一致时禁止继续"),
("策略漂移","实测指标持续超出历史置信范围时重新研究"),
],[1900,7460])
h2("第24周：结业答辩")
for x in ["10分钟说明策略假设和经济逻辑。","展示数据处理、回测、样本外和压力测试证据。","现场解释任意一笔交易为什么发生。","说明策略可能失效的三个具体场景。","展示模拟运行日志及其与回测假设的差异。"]: bullet(x)
callout("结业标准", "综合评分达到75分，且“回测正确性”和“风险纪律”两项均不低于70%；否则即使历史收益很高，也不视为结业。")

h1("学习成果检测体系")
h2("四层检测法")
table(["层级","要回答的问题","证据"],[
("记忆","我是否记得概念？","闭卷术语解释、公式默写"),
("理解","我是否能解释原因？","口述、反例、画流程图"),
("应用","我是否能独立完成？","代码、测试、图表、报告"),
("审计","我是否能发现错误？","单步回放、偏差检查、压力测试"),
],[1200,3800,4360])
h2("每周自测模板")
for x in ["本周我能不查资料讲清楚的三个概念是什么？","我独立写出的代码或报告是什么？","有哪些结果与预期不一致？为什么？","本周发现的一个错误是什么？如何防止再次出现？","下周进入新内容前，必须补齐什么？"]: bullet(x)
h2("综合评分表（100分）")
table(["维度","分值","评分依据"],[
("编程与数据","15","代码可读、数据校验充分、异常处理明确"),
("金融机制","10","成交、复权、费用和市场约束理解正确"),
("统计基础","10","指标与检验使用合理，不滥用显著性"),
("回测正确性","25","无未来泄漏，逐日核对和自动测试通过"),
("研究方法","15","假设清晰、样本外、稳定性与反证充分"),
("风险控制","15","仓位、回撤、集中度和故障规则明确"),
("可复现性","10","数据、代码、参数和环境可追溯"),
],[1900,900,6560])

h1("必做测试清单")
h2("数据测试")
for x in ["日期严格递增且无重复。","OHLC关系合理：high不低于open/close，low不高于open/close。","价格、成交量单位和复权方式明确。","缺失交易日有记录，不用未经解释的0填充。","多资产对齐规则明确，避免把停牌误当价格不变且可成交。"]: bullet(x)
h2("回测测试")
for x in ["用5～10日手工样本逐日核对。","将未来数据截断，过去结果不应改变。","信号整体延迟一期后结果符合预期。","费用设为0和较高值时，净值变化方向正确。","没有持仓时策略收益为0；满仓时与标的收益一致。","现金、持仓市值和总资产始终满足会计恒等关系。","重复运行结果完全一致。"]: bullet(x)
h2("研究测试")
for x in ["基准选择合理。","不只报告最佳参数。","测试期未用于调参。","报告失败年份和最差交易。","策略逻辑可用一句话说明。","若去掉收益最大的5笔交易，结果仍被讨论。"]: bullet(x)

h1("研究报告模板")
h2("1. 摘要")
para("一句话策略、研究资产和区间、核心结论、主要风险。摘要必须同时写优点和限制。")
h2("2. 假设与机制")
para("为什么该信号可能获得收益？谁提供风险溢价或行为偏差？在什么环境下应失效？")
h2("3. 数据")
para("来源、字段、频率、复权、时间范围、股票池定义、清洗规则、信息可得时间。")
h2("4. 规则")
para("信号公式、观察时间、下单时间、成交价格、仓位、退出、费用、滑点和约束。")
h2("5. 结果")
para("净值、回撤、年度/月度收益、交易明细、换手、风险暴露及基准对比。")
h2("6. 验证")
para("样本内外、Walk-forward、参数敏感性、成本压力、信号延迟、市场分段和反事实测试。")
h2("7. 局限与下一步")
para("列出数据、模型、成交和制度假设的限制；下一步实验必须事先定义，避免看到结果后任意调整。")

h1("学习日志与复盘表")
h2("每日学习记录")
table(["日期","主题","投入时间","输出物","卡点/下一步"],[
("____","________________","____小时","________________","________________"),
("____","________________","____小时","________________","________________"),
("____","________________","____小时","________________","________________"),
],[900,1700,1200,2600,2960])
h2("每周复盘")
table(["项目","填写内容"],[
("本周完成","____________________________________________________________"),
("最重要理解","____________________________________________________________"),
("发现的错误","____________________________________________________________"),
("验收结果","通过 / 有条件通过 / 未通过"),
("补救任务","____________________________________________________________"),
("下周目标","____________________________________________________________"),
],[1800,7560])
h2("策略变更记录")
table(["版本","日期","修改前规则","修改原因","是否重新样本外验证"],[
("v___","____","____________","____________","是 / 否"),
("v___","____","____________","____________","是 / 否"),
],[900,1000,2600,2860,2000])

h1("推荐学习资源与工具")
para("优先选择官方文档、经典教材和能提供完整代码与数据时点说明的课程。工具版本会变化，安装时应查阅对应官方文档。")
for x in [
"NumPy官方初学者指南：https://numpy.org/doc/stable/user/absolute_beginners.html",
"pandas官方入门教程：https://pandas.pydata.org/docs/getting_started/intro_tutorials/",
"Matplotlib官方教程：https://matplotlib.org/stable/tutorials/index.html",
"Backtrader快速入门：https://www.backtrader.com/docu/quickstart/quickstart/",
"QuantConnect LEAN文档：https://www.quantconnect.com/docs/v2/lean-engine/getting-started",
]: bullet(x)
h2("资源筛选标准")
for x in ["是否明确数据时点和成交时点。","是否计入交易成本和滑点。","是否有样本外结果，而非只展示最优回测。","是否提供失败案例和适用边界。","是否能复现，而不是只展示收益截图。"]: bullet(x)

h1("常见误区与纠偏")
table(["误区","为什么危险","正确动作"],[
("先找稳赚策略","忽略研究方法和风险","先复现简单基线并审计"),
("收益越高越好","可能来自杠杆、尾部风险或错误","同时看回撤、换手和稳定性"),
("参数越精细越专业","参数多会加剧过拟合","减少自由度，查看参数区域"),
("机器学习必然更强","复杂模型更易泄漏和过拟合","先与简单线性/规则基线比较"),
("回测好就实盘","执行、成本和数据延迟会改变结果","先持续模拟并核对偏差"),
("亏损就立即改规则","把随机波动当模型失效","按预先设定的评估窗口处理"),
],[1600,3800,3960])

h1("最终结业检查清单")
for x in [
"我能独立获取、清洗和验证一份时间序列数据。",
"我能解释每一列数据何时对交易者可见。",
"我能从零实现收益、回撤、波动率和夏普。",
"我能区分信号、仓位、交易和策略收益。",
"我能用小样本逐日核对回测。",
"我的回测包含合理费用、滑点和成交约束。",
"我能识别未来函数、幸存者偏差和过拟合。",
"我保留了真正未参与调参的测试集。",
"我完成了参数、成本、延迟和市场分段压力测试。",
"我能解释收益来源及三个可能失效场景。",
"我有完整的研究日志、版本记录和可复现说明。",
"我完成至少一个月的模拟运行，没有临时修改规则。",
]: bullet("□ "+x)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.keep_together = True
font(p.add_run("最后提醒："), 9.5, True, RED)
font(p.add_run("量化能力的标志不是画出漂亮净值，而是能怀疑结果、定位错误、量化风险并让别人复现研究。只有在理解风险、完成长期模拟并能承受损失时，才考虑真实资金。"), 9.5)

# Core properties
doc.core_properties.title = "量化交易与量化投资系统学习手册"
doc.core_properties.subject = "24周学习路线、内容、项目与成果验收"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "量化交易, 量化投资, Python, 回测, 风险管理, 学习路线"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
